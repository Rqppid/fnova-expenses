"""Read-side model of a tracker workbook: rows, totals, structure."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from .layout import (COL_AMOUNT, COL_DATE, COL_DESC, COL_RECEIPT, COL_REF, COL_REMARKS,
                     COL_VAT, COL_VATFLAG, COL_VENDOR, NET_LABEL, TOTAL_LABEL, Layout)

DATE_RE = re.compile(r"^(\d{1,2})\.(\d{1,2})\.(\d{2})$")
SUMIFS_RANGE_RE = re.compile(r"(Sheet2!\$[A-Z]{1,2}\$\d+:\$[A-Z]{1,2}\$)(\d+)")


class TrackerError(Exception):
    pass


def parse_ddmmyy(text) -> date | None:
    m = DATE_RE.match(str(text or "").strip())
    if not m:
        return None
    d, mo, y = (int(g) for g in m.groups())
    try:
        return date(2000 + y, mo, d)
    except ValueError:
        return None


def parse_amount(value) -> float | None:
    """Numeric cell value, or a hand-typed text amount like '£12.50'. None if unparseable."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        s = value.strip().replace("£", "").replace(",", "").replace("GBP", "").strip()
        try:
            return float(s)
        except ValueError:
            return None
    return None


@dataclass
class Row:
    r: int
    date_text: str | None
    date: date | None
    vendor: str
    desc: str
    amount_raw: object
    amount: float | None
    ref: str
    receipt: str
    vat_flag: str
    vat_raw: object
    vat: float | None
    remarks: str
    helpers: dict[str, object] = field(default_factory=dict)

    @property
    def amount_is_numeric(self) -> bool:
        return _is_num(self.amount_raw)


def find_total_row(ws: Worksheet, layout: Layout) -> int:
    """Row holding 'Total :' in column D. Net Total must sit directly below."""
    hits = [c.row for c in ws[COL_DATE]
            if isinstance(c.value, str) and c.value.strip() == TOTAL_LABEL.strip()
            and c.row >= layout.first_row]
    if len(hits) != 1:
        raise TrackerError(f"Expected exactly one '{TOTAL_LABEL}' row in column D, found {hits}")
    t = hits[0]
    net = ws[f"{COL_DATE}{t + 1}"].value
    if not (isinstance(net, str) and net.strip().startswith("Net Total")):
        raise TrackerError(f"Row {t + 1} should be '{NET_LABEL}', found {net!r}")
    return t


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _s(v) -> str:
    return "" if v is None else str(v)


def read_row(ws: Worksheet, layout: Layout, r: int) -> Row:
    g = lambda col: ws[f"{col}{r}"].value  # noqa: E731
    dt = g(COL_DATE)
    return Row(
        r=r, date_text=_s(dt) or None, date=parse_ddmmyy(dt),
        vendor=_s(g(COL_VENDOR)), desc=_s(g(COL_DESC)),
        amount_raw=g(COL_AMOUNT), amount=parse_amount(g(COL_AMOUNT)),
        ref=_s(g(COL_REF)), receipt=_s(g(COL_RECEIPT)), vat_flag=_s(g(COL_VATFLAG)).strip(),
        vat_raw=g(COL_VAT), vat=parse_amount(g(COL_VAT)), remarks=_s(g(COL_REMARKS)).strip(),
        helpers={h: g(h) for h in layout.helpers},
    )


@dataclass
class Snapshot:
    layout: Layout
    total_row: int
    rows: list[Row]
    total_formulas: dict[str, object]
    net_formulas: dict[str, object]
    sumifs: dict[str, str]          # coordinate -> formula on summary sheet

    @property
    def last_row(self) -> int:
        return self.total_row - 1

    @property
    def gross(self) -> float:
        # Mirrors Excel's SUM, which ignores text cells such as a hand-typed '£12.50'.
        return round(sum(r.amount for r in self.rows if r.amount_is_numeric), 2)

    @property
    def vat(self) -> float:
        return round(sum(r.vat_raw for r in self.rows if _is_num(r.vat_raw)), 2)

    @property
    def net(self) -> float:
        return round(self.gross - self.vat, 2)

    def monthly(self) -> dict[tuple[int, int], tuple[float, float]]:
        out: dict[tuple[int, int], list[float]] = {}
        for r in self.rows:
            if r.date is None:
                continue
            k = (r.date.year, r.date.month)
            acc = out.setdefault(k, [0.0, 0.0])
            acc[0] += r.amount if r.amount_is_numeric else 0.0
            acc[1] += r.vat_raw if _is_num(r.vat_raw) else 0.0
        return {k: (round(a, 2), round(b, 2)) for k, (a, b) in sorted(out.items())}


def read_snapshot_ws(wb, layout: Layout) -> Snapshot:
    ws = wb[layout.log_sheet]
    t = find_total_row(ws, layout)
    rows = [read_row(ws, layout, r) for r in range(layout.first_row, t)]
    cols = ("D", "I", "Q")
    total = {c: ws[f"{c}{t}"].value for c in cols}
    net = {c: ws[f"{c}{t + 1}"].value for c in cols}
    s1 = wb[layout.summary_sheet]
    sumifs = {c.coordinate: c.value for row in s1.iter_rows() for c in row
              if isinstance(c.value, str) and "SUMIFS(" in c.value}
    return Snapshot(layout, t, rows, total, net, sumifs)


def read_snapshot(path: Path, layout: Layout) -> Snapshot:
    wb = openpyxl.load_workbook(path)
    return read_snapshot_ws(wb, layout)


def sumifs_end_rows(formula: str) -> set[int]:
    return {int(m.group(2)) for m in SUMIFS_RANGE_RE.finditer(formula)}
