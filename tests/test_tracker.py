import openpyxl
import pytest

from fonenova import tracker
from fonenova.sheet import TrackerError, find_total_row, read_snapshot
from fonenova.verify import verify
from fonenova.xlsx_diff import diff

BASE_GROSS, BASE_VAT, BASE_LAST = 8951.59, 423.06, 70


def test_reference_state(septdec_copy, lays):
    s = read_snapshot(septdec_copy, lays["septdec"])
    assert (s.last_row, s.total_row) == (BASE_LAST, 71)
    assert (s.gross, s.vat, s.net) == (BASE_GROSS, BASE_VAT, 8528.53)
    assert verify(septdec_copy, lays["septdec"]).ok


def test_add_row(septdec_copy, lays, bdir):
    lay = lays["septdec"]
    res = tracker.add_row(lay, bdir, date="06.10.26", vendor="Test Vendor", desc="Test",
                          amount=12.0, ref="T1", vat_flag="VAT-Yes", vat_amount=2.0,
                          remarks="test row", path=septdec_copy)
    assert (res.last_row, res.gross, res.vat) == (71, round(BASE_GROSS + 12, 2), round(BASE_VAT + 2, 2))
    wb = openpyxl.load_workbook(septdec_copy)
    ws = wb["Sheet2"]
    assert ws["D71"].value == "06.10.26" and ws["I71"].value == 12.0
    assert ws["D72"].value == "Total :" and ws["I72"].value == "=SUM(I5:I71)"
    assert ws["I73"].value == "=I72-SUM(Q5:Q71)"
    assert ws["U71"].value.startswith("=IFERROR(DATE(") and "D71" in ws["U71"].value
    assert ws["I71"].style_id == ws["I70"].style_id
    for c in ("B5", "C5", "B8", "C8"):
        assert "$I$71" in wb["Sheet1"][c].value or "$Q$71" in wb["Sheet1"][c].value
        assert "$U$71" in wb["Sheet1"][c].value
    assert (bdir / next(bdir.iterdir()).name).exists()


def test_add_twice_finds_moved_total(septdec_copy, lays, bdir):
    lay = lays["septdec"]
    for i in range(2):
        tracker.add_row(lay, bdir, date="06.10.26", vendor=f"V{i}", desc="d", amount=1.0, ref="r",
                        vat_flag="VAT-No", vat_amount=None, remarks="x", path=septdec_copy)
    s = read_snapshot(septdec_copy, lay)
    assert (s.last_row, s.total_row) == (72, 73)
    assert verify(septdec_copy, lay).ok


def test_add_preserves_features(septdec_copy, lays, bdir, tmp_path):
    import shutil
    before = tmp_path / "before.xlsx"
    shutil.copy2(septdec_copy, before)
    tracker.add_row(lays["septdec"], bdir, date="06.10.26", vendor="V", desc="d", amount=1.0,
                    ref="r", vat_flag="VAT-No", vat_amount=None, remarks="x", path=septdec_copy)
    d = diff(before, septdec_copy)
    assert d["features"] == []
    changed = {c.split(":")[0].split("!")[1] for c in d["cells"]}
    # Only the new row, the moved totals (rows 71-73, any column's style) and the 8 SUMIFS change.
    import re
    assert all(re.fullmatch(r"[A-Z]+7[123]", c) or re.fullmatch(r"[BC][5-8]", c) for c in changed), changed


def test_validation_errors(septdec_copy, lays, bdir):
    lay = lays["septdec"]
    kw = dict(vendor="V", desc="d", ref="r", remarks="x", path=septdec_copy)
    with pytest.raises(TrackerError):
        tracker.add_row(lay, bdir, date="6/10/2026", amount=1.0, vat_flag="VAT-No", vat_amount=None, **kw)
    with pytest.raises(TrackerError):
        tracker.add_row(lay, bdir, date="06.10.26", amount=1.0, vat_flag="VAT-Yes", vat_amount=None, **kw)
    with pytest.raises(TrackerError):
        tracker.add_row(lay, bdir, date="06.10.26", amount=1.0, vat_flag="Yes", vat_amount=None, **kw)


def test_refuses_when_excel_open(septdec_copy, lays, bdir):
    (septdec_copy.parent / ("~$" + septdec_copy.name)).write_bytes(b"x")
    with pytest.raises(TrackerError, match="Excel has the file open"):
        tracker.add_row(lays["septdec"], bdir, date="06.10.26", vendor="V", desc="d", amount=1.0,
                        ref="r", vat_flag="VAT-No", vat_amount=None, remarks="x", path=septdec_copy)


def test_v18_is_read_only(v18_ref, lays, bdir):
    with pytest.raises(TrackerError, match="read-only"):
        tracker.add_row(lays["v18"], bdir, date="06.09.26", vendor="V", desc="d", amount=1.0,
                        ref="r", vat_flag="VAT-No", vat_amount=None, remarks="x", path=v18_ref)


def test_delete_row(septdec_copy, lays, bdir):
    lay = lays["septdec"]
    wb = openpyxl.load_workbook(septdec_copy)
    row6 = [c.value for c in wb["Sheet2"][6]][3:18]
    res = tracker.delete_row(lay, bdir, row=5, expect_date="01.09.26", expect_vendor="Tesco",
                             expect_amount=5.5, path=septdec_copy)
    assert (res.last_row, res.gross, res.vat) == (69, round(BASE_GROSS - 5.5, 2), BASE_VAT)
    ws = openpyxl.load_workbook(septdec_copy)["Sheet2"]
    assert [c.value for c in ws[5]][3:18] == row6          # row 6 moved up intact
    assert ws["D70"].value == "Total :" and ws["I70"].value == "=SUM(I5:I69)"
    assert ws["I71"].value == "=I70-SUM(Q5:Q69)"
    assert all(c.value is None for c in ws[72])
    assert "D5" in ws["U5"].value


def test_delete_guard(septdec_copy, lays, bdir):
    with pytest.raises(TrackerError, match="not"):
        tracker.delete_row(lays["septdec"], bdir, row=5, expect_date="01.09.26", expect_vendor="Lidl",
                           expect_amount=5.5, path=septdec_copy)


def test_complete_half_finished_row(septdec_copy, lays, bdir):
    lay = lays["septdec"]
    # Simulate a hand-typed row: text amount, no helpers, no Remarks.
    tracker.add_row(lay, bdir, date="06.10.26", vendor="Hand", desc="d", amount=1.0, ref="r",
                    vat_flag="VAT-No", vat_amount=None, remarks="x", path=septdec_copy)
    wb = openpyxl.load_workbook(septdec_copy)
    ws = wb["Sheet2"]
    ws["I71"], ws["R71"], ws["U71"], ws["V71"], ws["W71"] = "£12.50", None, None, None, None
    wb.save(septdec_copy)
    assert not verify(septdec_copy, lay).ok
    res = tracker.complete_row(lay, bdir, row=71, remarks="completed", path=septdec_copy)
    assert res.gross == round(BASE_GROSS + 12.5, 2)
    assert verify(septdec_copy, lay).ok


def test_verify_catches_short_sumifs(septdec_copy, lays):
    wb = openpyxl.load_workbook(septdec_copy)
    wb["Sheet1"]["B5"] = wb["Sheet1"]["B5"].value.replace("$70", "$69")
    wb.save(septdec_copy)
    res = verify(septdec_copy, lays["septdec"])
    assert not res.ok and any("B5" in p for p in res.problems)


def test_v18_reads(v18_ref, lays):
    s = read_snapshot(v18_ref, lays["v18"])
    assert (s.last_row, s.gross, s.vat) == (313, 34370.55, 1911.59)
