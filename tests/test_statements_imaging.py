from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest
from PIL import Image, ImageDraw

from fonenova import statements as S
from fonenova.filing import file_receipt_pdf
from fonenova.sheet import read_snapshot
from fonenova.tracker import add_rows, fx_entries
from fonenova.verify import verify

HEADER = ("Date started (UTC),Date completed (UTC),ID,Type,State,Description,Reference,Payer,Card number,"
          "Card label,Card state,Orig currency,Orig amount,Payment currency,Amount,Total amount,"
          "Exchange rate,Fee,Fee currency,Balance,Account\n")


def _csv(tmp_path, lines):
    p = tmp_path / "statement.csv"
    p.write_text(HEADER + "\n".join(lines) + "\n", encoding="utf-8")
    return p


def _row(d, i, typ, desc, amt, ref="", orig_cur="GBP", orig_amt=None, rate="", fee="0.00"):
    orig_amt = orig_amt if orig_amt is not None else abs(amt)
    return (f"{d},{d},{i},{typ},COMPLETED,{desc},{ref},W Ismati,516760******9873,Standard,ACTIVE,"
            f"{orig_cur},{orig_amt},GBP,{amt},{amt},{rate},{fee},GBP,100.00,GBP Main")


@pytest.fixture
def setup(lays, septdec_copy, v18_ref):
    live = replace(lays["septdec"], path=septdec_copy)
    v18 = replace(lays["v18"], path=v18_ref)
    snaps = {"septdec": read_snapshot(septdec_copy, live), "v18": read_snapshot(v18_ref, v18)}
    return live, snaps


def test_reconcile_classifies(tmp_path, setup):
    live, snaps = setup
    csvp = _csv(tmp_path, [
        _row("2026-10-04", "aaaaaaaa-1111-0000", "CARD_PAYMENT", "Tesco Stores 4453", -6.00),       # row exists
        _row("2026-10-06", "bbbbbbbb-2222-0000", "CARD_PAYMENT", "Costa Coffee 1234", -4.20),       # new
        _row("2026-10-02", "cccccccc-3333-0000", "FEE", "Revolut Business Fee", -10.00, "Basic plan fee"),
        _row("2026-10-03", "dddddddd-4444-0000", "TRANSFER", "To Mobile One Trade", -900.00),       # wholesale
        _row("2026-10-03", "eeeeeeee-5555-0000", "TRANSFER", "To Someone", -50.00),                 # transfer
        _row("2026-08-20", "ffffffff-6666-0000", "CARD_PAYMENT", "Old Shop", -9.99),                # V18 period
        _row("2026-10-05", "99999999-7777-0000", "EXCHANGE", "Main EUR -> Main GBP", 850.00, "",
             "EUR", -1000.00, "1.176471", "-3.50"),
        _row("2026-10-05", "88888888-8888-0000", "TOPUP", "Money added", 500.00),
    ])
    rec = S.reconcile(S.parse_csv(csvp), snaps, live)
    assert [t.description for t in rec.to_log] == ["Revolut Business Fee", "Costa Coffee 1234"]
    assert any("Tesco" in t.description for t, _ in rec.matched)
    reasons = {t.description: why for t, why in rec.review}
    assert "wholesale" in reasons["To Mobile One Trade"] and "transfer" in reasons["To Someone"]
    assert [t.description for t in rec.v18_period] == ["Old Shop"]
    rows, fx = S.rows_for(rec, "statement.csv")
    assert [r["vendor"] for r in rows] == ["Revolut Business Fee", "Costa Coffee", "Revolut"]
    assert all(r["vat_flag"] == "VAT-No" for r in rows)
    assert "RECEIPT MISSING" in rows[1]["remarks"] and rows[2]["amount"] == 3.50
    assert fx == [dict(date="05.10.26", eur=1000.0, gbp=850.0, rate=1.176471, fee=3.5,
                       note="Revolut exchange Main EUR -> Main GBP, txn 99999999-7777-0000 (statement.csv)")]


def test_log_then_relog_is_idempotent(tmp_path, setup, bdir):
    live, snaps = setup
    csvp = _csv(tmp_path, [_row("2026-10-06", "bbbbbbbb-2222-0000", "CARD_PAYMENT", "Costa Coffee 1234", -4.20),
                           _row("2026-10-05", "99999999-7777-0000", "EXCHANGE", "EUR to GBP", 850.00, "",
                                "EUR", -1000.00, "1.176471", "-3.50")])
    before = snaps["septdec"]
    rec = S.reconcile(S.parse_csv(csvp), snaps, live)
    rows, fx = S.rows_for(rec, "statement.csv", fx_entries(openpyxl.load_workbook(live.path)))
    res = add_rows(live, bdir, rows, fx=fx, label="statement", path=live.path)
    assert res.last_row == before.last_row + 2 and res.gross == round(before.gross + 7.70, 2)
    assert res.vat == before.vat and verify(live.path, live).ok
    ws = openpyxl.load_workbook(live.path)["FX Exchanges (Reference)"]
    assert ws["A5"].value == "05.10.26" and ws["B5"].value == 1000.0 and ws["C5"].value == 850.0
    # Second run of the same statement: everything matches by txn id, nothing to add.
    snaps2 = {"septdec": read_snapshot(live.path, live), "v18": snaps["v18"]}
    rec2 = S.reconcile(S.parse_csv(csvp), snaps2, live)
    rows2, fx2 = S.rows_for(rec2, "statement.csv", fx_entries(openpyxl.load_workbook(live.path)))
    assert rows2 == [] and fx2 == [] and not rec2.to_log


def test_fx_confirmation_pdf_text():
    text = ("Transaction of 19 Sept 2026\nDescription Status Money out Money in\n"
            "Main · EUR -> Main · GBP Completed €4 670.00\nFX Rate EUR 1 = GBP 0.857067, Fee: £58.04 £3 944.46")
    assert S.parse_fx_confirmation(text) == {"date": "19.09.26", "eur": 4670.0, "gbp": 3944.46,
                                             "rate": 1.16677, "fee": 58.04}


def _photo(path: Path, rotate_exif=False):
    img = Image.new("RGB", (900, 1400), (60, 60, 60))                  # dark table
    d = ImageDraw.Draw(img)
    d.rectangle([150, 100, 750, 1300], fill=(235, 232, 225))            # receipt paper
    for i in range(20):
        d.text((190, 140 + i * 50), f"ITEM {i}  £{i}.99", fill=(120, 120, 120))
    d.text((190, 1220), "VAT 20%  £3.33   TOTAL £19.99", fill=(150, 150, 150))
    img.save(path, quality=85)
    return path


def test_photo_becomes_greyscale_pdf_and_original_is_kept(tmp_path):
    root = tmp_path / "receipts"
    (root / "Tesco").mkdir(parents=True)
    a, b = _photo(tmp_path / "IMG_1.jpg"), _photo(tmp_path / "IMG_2.jpg")
    processed = tmp_path / "Receipts Inbox" / "processed" / "2026-10-06"
    r = file_receipt_pdf([a, b], root, processed, vendor="Tesco Express", date="06.10.26")
    assert r.dest == root / "Tesco" / "Tesco-Express-06.10.26.pdf"
    import pdfplumber
    with pdfplumber.open(r.dest) as pdf:
        assert len(pdf.pages) == 2                                   # long receipt, two photos
        page = pdf.pages[0]
        assert page.width < page.height                              # cropped to the paper strip
        im = page.images[0]
        assert im["colorspace"][0] in ("DeviceGray", "/DeviceGray") or "Gray" in str(im["colorspace"])
    assert r.dest.stat().st_size < 500_000
    assert not a.exists() and (processed / "IMG_1.jpg").exists() and (processed / "IMG_2.jpg").exists()


def test_pdf_source_is_filed_unchanged(tmp_path):
    root = tmp_path / "receipts"
    root.mkdir()
    src = tmp_path / "invoice.pdf"
    src.write_bytes(b"%PDF-1.4 invoice")
    r = file_receipt_pdf([src], root, tmp_path / "processed", vendor="Xero UK", date="06.10.26")
    assert r.dest.read_bytes() == b"%PDF-1.4 invoice" and r.dest.suffix == ".pdf"
